r"""
Downloads every file from this term's courses and pulls the text out of them.

    python -m pip install pymupdf python-pptx python-docx
    python download.py

Files land in _originals\, extracted text in extracted\.
Run from inside the scraper folder.
"""

import json
import re
import sys
from pathlib import Path

from collect import BASE, LE, current_courses, get, get_client, walk_content

import paths

HERE = Path(__file__).parent
ORIGINALS = paths.ORIGINALS
EXTRACTED = paths.EXTRACTED

# Text pullers are optional; report what's missing instead of crashing.
READERS = {}
try:
    import pymupdf
    READERS["pdf"] = lambda p: "\n".join(pg.get_text() for pg in pymupdf.open(p))
except ImportError:
    try:
        from pypdf import PdfReader
        READERS["pdf"] = lambda p: "\n".join(
            (pg.extract_text() or "") for pg in PdfReader(p).pages
        )
    except ImportError:
        pass

try:
    from pptx import Presentation

    def _pptx(path):
        out = []
        for i, slide in enumerate(Presentation(path).slides, 1):
            out.append(f"\n--- slide {i} ---")
            for shape in slide.shapes:
                if shape.has_text_frame:
                    out.append(shape.text_frame.text)
            if slide.has_notes_slide:
                notes = slide.notes_slide.notes_text_frame.text.strip()
                if notes:
                    out.append(f"[speaker notes] {notes}")
        return "\n".join(out)

    READERS["pptx"] = _pptx
except ImportError:
    pass

try:
    import docx

    READERS["docx"] = lambda p: "\n".join(
        para.text for para in docx.Document(p).paragraphs
    )
except ImportError:
    pass

try:
    import openpyxl

    def _xlsx(path):
        out = []
        for sheet in openpyxl.load_workbook(path, data_only=True).worksheets:
            out.append(f"\n--- sheet: {sheet.title} ---")
            for row in sheet.iter_rows(values_only=True):
                cells = [str(c) for c in row if c is not None]
                if cells:
                    out.append("\t".join(cells))
        return "\n".join(out)

    READERS["xlsx"] = _xlsx
except ImportError:
    pass

READERS["txt"] = lambda p: Path(p).read_text(encoding="utf-8", errors="replace")


def safe_name(text, limit=80):
    cleaned = re.sub(r'[<>:"/\\|?*\x00-\x1f]', "_", (text or "untitled").strip())
    return (cleaned[:limit] or "untitled").rstrip(". ")


def extension(url, title):
    for candidate in (url or "", title or ""):
        match = re.search(r"\.([A-Za-z0-9]{1,5})(?:\?|$)", candidate)
        if match:
            return match.group(1).lower()
    return ""


def download_topic(client, topic, folder):
    """Fetch one file. Returns (path, note) -- path is None if skipped."""
    url = topic.get("url")
    if not url:
        return None, "no link"

    kind = (topic.get("type") or "").lower()
    if kind and kind not in ("file", "contentservice"):
        return None, f"not a file ({kind})"

    full = url if url.startswith("http") else BASE + ("" if url.startswith("/") else "/") + url
    ext = extension(url, topic.get("title"))
    name = safe_name(topic.get("title")) + (f".{ext}" if ext else "")
    dest = folder / name

    if dest.exists() and dest.stat().st_size > 0:
        return dest, "already had it"

    try:
        r = client.get(full)
    except Exception as e:
        return None, f"failed: {e!r}"

    if r.status_code != 200:
        return None, f"failed: HTTP {r.status_code}"
    if b"<html" in r.content[:400].lower() and ext not in ("html", "htm"):
        return None, "got a web page, not a file"

    folder.mkdir(parents=True, exist_ok=True)
    dest.write_bytes(r.content)
    return dest, f"{len(r.content) // 1024} KB"


def extract(path, out_dir):
    ext = path.suffix.lstrip(".").lower()
    reader = READERS.get(ext)
    if not reader:
        return None, f"can't read .{ext}"
    try:
        text = reader(str(path))
    except Exception as e:
        return None, f"unreadable: {e!r}"

    text = re.sub(r"\n{3,}", "\n\n", (text or "").strip())
    if len(text) < 20:
        return None, "no text (probably a scan)"

    out_dir.mkdir(parents=True, exist_ok=True)
    out = out_dir / (path.stem + ".txt")
    out.write_text(text, encoding="utf-8")
    return out, f"{len(text.split()):,} words"


# Where a file can hang besides Content, and the endpoint that serves it.
# collect.py gathers twelve tabs; this used to download from one. GNG2101's
# ten deliverable briefs, every template, every lab manual and the project
# list attached to an announcement all live here -- 49 files and instruction
# blocks the app had never read.
#
# The dropbox URL was confirmed against the real thing before this was
# written (probe_attachments.py --live: HTTP 200, right content type, exact
# byte count). The news one is not confirmed, so a failure has to be loud
# rather than silent -- that is the whole point of this change.
ATTACHMENT_SOURCES = [
    ("assignments", "Name", "Id",
     "/d2l/api/le/{le}/{oid}/dropbox/folders/{item}/attachments/{file}"),
    ("announcements", "Title", "Id",
     "/d2l/api/le/{le}/{oid}/news/{item}/attachments/{file}"),
]


def instruction_text(value):
    """A folder's typed instructions as plain text, or ''. Shapes vary."""
    if isinstance(value, dict):
        value = value.get("Text") or value.get("Html") or ""
    return re.sub(r"<[^>]+>", " ", str(value or "")).strip()


def download_attachments(client, raw, code):
    """Files hanging off tabs other than Content, plus typed instructions.

    `raw` is this course's entry from collected.json -- collect.py has
    already fetched these lists, so nothing here re-asks for them.

    Returns (downloaded, readable, words, failures).
    """
    got = read = words = 0
    failures = []
    oid = raw.get("id")
    originals, extracted = ORIGINALS / code, EXTRACTED / code

    for tab, title_key, id_key, template in ATTACHMENT_SOURCES:
        for item in raw.get(tab) or []:
            if not isinstance(item, dict):
                continue
            item_id = item.get(id_key)
            for f in item.get("Attachments") or []:
                if not isinstance(f, dict):
                    continue
                file_id = f.get("FileId") or f.get("Id")
                name = safe_name(f.get("FileName") or f.get("Name") or "file")
                if not (item_id and file_id):
                    continue
                dest = originals / name
                label = name[:46]

                if dest.exists() and dest.stat().st_size > 0:
                    out = extracted / (dest.stem + ".txt")
                    if out.exists():
                        continue                       # had it, already read
                else:
                    url = template.format(le=LE, oid=oid, item=item_id, file=file_id)
                    try:
                        r = client.get(BASE + url, follow_redirects=True)
                    except Exception as exc:
                        failures.append((tab, name, f"{type(exc).__name__}"))
                        print(f"  FAIL  {label:<48} {type(exc).__name__}")
                        continue
                    if r.status_code != 200:
                        failures.append((tab, name, f"HTTP {r.status_code}"))
                        print(f"  FAIL  {label:<48} HTTP {r.status_code}  [{tab}]")
                        continue
                    originals.mkdir(parents=True, exist_ok=True)
                    dest.write_bytes(r.content)
                    got += 1

                out, detail = extract(dest, extracted)
                if out:
                    read += 1
                    words += int(detail.split()[0].replace(",", ""))
                    print(f"  read  {label:<48} {detail}  [{tab}]")
                else:
                    print(f"  saved {label:<48} {detail}  [{tab}]")

    # Typed instructions are already in hand -- no request, no failure mode,
    # and on GNG2101 they are 553 characters of what Deliverable A asks for.
    for item in raw.get("assignments") or []:
        if not isinstance(item, dict):
            continue
        body = instruction_text(item.get("CustomInstructions"))
        if len(body) < 20:
            continue
        title = safe_name(item.get("Name") or "assignment", 60)
        out = extracted / f"{title} (instructions).txt"
        text = f"{item.get('Name') or ''}\n\n{body}\n"
        if out.exists() and out.read_text(encoding="utf-8") == text:
            continue
        extracted.mkdir(parents=True, exist_ok=True)
        out.write_text(text, encoding="utf-8")
        read += 1
        words += len(body.split())
        print(f"  read  {title[:46]:<48} {len(body.split()):,} words  "
              f"[instructions]")

    return got, read, words, failures


def main(client=None):
    if "pdf" not in READERS:
        print("No PDF reader installed. Run:")
        print("    python -m pip install pymupdf python-pptx python-docx")
        print("If pymupdf will not install on Python 3.14, use pypdf instead:")
        print("    python -m pip install pypdf python-pptx python-docx\n")

    client = client or get_client()
    courses, _ = current_courses(client)

    grand_files = grand_text = 0
    words_total = 0
    all_failures = []

    # collect.py has already fetched every tab, so the attachment lists are
    # in hand and none of this costs a request. Run standalone after a long
    # gap and this is simply the last scrape's copy.
    raw_by_id = {}
    if paths.COLLECTED.exists():
        try:
            raw_by_id = {c.get("id"): c for c in
                         json.loads(paths.COLLECTED.read_text(encoding="utf-8"))}
        except (json.JSONDecodeError, OSError) as exc:
            print(f"  (could not read collected.json: {exc} -- "
                  "attachments skipped this run)")

    for course in courses:
        code = safe_name(course["name"], 40)
        print(f"\n{course['name']}")
        _, topics = walk_content(client, course["id"])
        if not topics:
            print("  (no files posted yet)")
            continue

        got = read = 0
        links = []
        for topic in topics:
            path, note = download_topic(client, topic, ORIGINALS / code)
            label = (topic.get("title") or "?")[:46]
            if not path:
                print(f"  skip  {label:<48} {note}")
                if note.startswith("not a file"):
                    links.append(topic)
                continue
            got += 1
            out, detail = extract(path, EXTRACTED / code)
            if out:
                read += 1
                words_total += int(detail.split()[0].replace(",", ""))
                print(f"  read  {label:<48} {detail}")
            else:
                print(f"  saved {label:<48} {detail}")

        if links:
            # These are usually assignment dropboxes or quizzes, which the
            # assignments and quizzes endpoints already report with real due
            # dates. Recorded rather than discarded so that can be checked.
            folder = EXTRACTED / code
            folder.mkdir(parents=True, exist_ok=True)
            (folder / "_links.txt").write_text(
                "\n".join(f"{t.get('title','?')}\t{t.get('url','')}" for t in links),
                encoding="utf-8",
            )

        raw = raw_by_id.get(course["id"])
        if raw:
            a_got, a_read, a_words, fails = download_attachments(client, raw, code)
            got += a_got
            read += a_read
            words_total += a_words
            all_failures += [(course["name"], *f) for f in fails]

        grand_files += got
        grand_text += read
        note = f", {len(links)} links noted" if links else ""
        print(f"  -> {got} downloaded, {read} readable{note}")

    print("\n" + "=" * 64)
    print(f"{grand_files} files downloaded, {grand_text} turned into readable text")
    print(f"about {words_total:,} words to search for deadlines")
    if all_failures:
        print(f"\n{len(all_failures)} attachment(s) could not be fetched -- "
              "these are files the app does not have:")
        for name, tab, fname, why in all_failures:
            print(f"    {why:<16} [{tab}]  {str(name)[:22]:<22} {fname[:40]}")

    print(f"\nOriginals: {ORIGINALS}")
    print(f"Text:      {EXTRACTED}")


if __name__ == "__main__":
    main()
